"""
Bulk Upload Service
Shared functionality for processing Excel bulk uploads for Customers, Employees, Services, and Products
"""
from io import BytesIO
import logging
from datetime import datetime
from decimal import Decimal, InvalidOperation
from flask import g

logger = logging.getLogger(__name__)

# Field mapping configurations for each module
FIELD_MAPPINGS = {
    'customers': {
        'First Name': 'first_name',
        'Last Name': 'last_name',
        'Phone Number': 'phone',
        'Email': 'email',
        'Gender': 'gender',
        'Date of Birth': 'date_of_birth',
        'Address': 'address',
        'Notes': 'notes',
        'Spot': 'spot'
    },
    'employees': {
        'First Name': 'first_name',
        'Last Name': 'last_name',
        'Phone Number': 'phone',
        'Specialization': 'specialization',
        'Role': 'role',
        'Salary': 'salary',
        'Commission %': 'commission_percentage',
        'Joining Date': 'joining_date',
        'Status': 'status'
    },
    'services': {
        'Service Name': 'name',
        'Category Name': 'category_name',
        'Price': 'price',
        'Duration (minutes)': 'duration_minutes',
        'Description': 'description',
        'Status': 'status',
        'Membership Plan Name': 'membership_plan_name',
        'Membership Discount (%)': 'membership_discount_percentage',
        'Membership Discount Amount': 'membership_discount_amount'
    },
    'products': {
        'Product Name': 'name',
        'Category': 'category',
        'SKU': 'sku',
        'Barcode': 'barcode',
        'Cost Price': 'cost_price',
        'Selling Price': 'selling_price',
        'MRP': 'mrp',
        'Stock Quantity': 'stock_quantity',
        'Low Stock Threshold': 'low_stock_threshold',
        'Status': 'status'
    }
}

# Required fields for each module
REQUIRED_FIELDS = {
    'customers': ['First Name', 'Phone Number'],
    'employees': ['First Name', 'Phone Number', 'Salary'],
    'services': ['Service Name', 'Category Name', 'Price'],
    'products': ['Product Name', 'Cost Price', 'Selling Price', 'MRP', 'Stock Quantity']
}

# Optional fields for each module
OPTIONAL_FIELDS = {
    'customers': ['Last Name', 'Email', 'Gender', 'Date of Birth', 'Address', 'Notes', 'Spot'],
    'employees': ['Last Name', 'Specialization', 'Role', 'Commission %', 'Joining Date', 'Status'],
    'services': ['Description', 'Status', 'Membership Plan Name', 'Membership Discount (%)', 'Membership Discount Amount'],
    'products': ['Category', 'SKU', 'Barcode', 'Low Stock Threshold', 'Status']
}

def get_membership_mode(tenant_id, branch_id=None):
    if not tenant_id:
        return "paid_plan"
    try:
        from app.models.visit_membership import VisitMembershipSetting
        from app.utils.auth import get_tenant_query
        was_master = getattr(g, 'use_master_db', False)
        g.use_master_db = False
        setting = get_tenant_query(VisitMembershipSetting).filter_by(tenant_id=tenant_id).first()
        g.use_master_db = was_master
        if setting and setting.membership_mode:
            return setting.membership_mode
    except Exception:
        pass
    return "paid_plan"

def generate_excel_template(module_name, tenant_id=None, branch_id=None):
    """
    Generate an Excel template for the specified module
    Returns the Excel file as bytes
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{module_name.capitalize()} Template"
        
        # Get field mapping for this module
        field_mapping = FIELD_MAPPINGS.get(module_name, {})
        mode = get_membership_mode(tenant_id, branch_id)

        if module_name == 'services' and mode not in ['paid_plan', 'discount_plan']:
            headers = ['Service Name', 'Category Name', 'Price', 'Duration (minutes)', 'Description', 'Status']
        else:
            headers = list(field_mapping.keys())
        
        # Style the header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="EC4899", end_color="EC4899", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Write headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Add sample data row
        sample_data = get_sample_data(module_name, membership_mode=mode)
        if sample_data:
            for row_num, row_data in enumerate(sample_data, 2):
                for col_num, header in enumerate(headers, 1):
                    field_key = field_mapping[header]
                    value = row_data.get(field_key, '')
                    ws.cell(row=row_num, column=col_num, value=value)
                
                # Pre-fill dynamic Excel formula for Membership Discount Amount if present
                if module_name == 'services' and 'Membership Discount Amount' in headers and 'Price' in headers and 'Membership Discount (%)' in headers:
                    price_col = openpyxl.utils.get_column_letter(headers.index('Price') + 1)
                    pct_col = openpyxl.utils.get_column_letter(headers.index('Membership Discount (%)') + 1)
                    amt_col_idx = headers.index('Membership Discount Amount') + 1
                    formula_str = f"=IF(AND(ISNUMBER({price_col}{row_num}), ISNUMBER({pct_col}{row_num})), ROUND({price_col}{row_num}*({pct_col}{row_num}/100), 2), \"\")"
                    ws.cell(row=row_num, column=amt_col_idx, value=formula_str)
        
        # Add Excel Data Validation Dropdowns
        from openpyxl.worksheet.datavalidation import DataValidation
        from app.utils.auth import get_tenant_query
        
        if module_name == 'services' and tenant_id:
            from app.models.membership import MembershipPlan
            from app.models.catalog import ServiceCategory
            was_master = getattr(g, 'use_master_db', False)
            g.use_master_db = False

            # Category Name Dropdown (Col B)
            try:
                cats = get_tenant_query(ServiceCategory).filter_by(tenant_id=tenant_id, is_deleted=False).all()
                cat_names = [c.name.replace('"', '').strip() for c in cats if c.name]
                if cat_names:
                    cat_formula = f'"{",".join(cat_names)}"'
                    dv_cat = DataValidation(type="list", formula1=cat_formula, allow_blank=True)
                    dv_cat.hide_drop_down = None
                    dv_cat.showInputMessage = True
                    dv_cat.promptTitle = "Category Selection"
                    dv_cat.prompt = "Select an existing category from the dropdown or type a new category name."
                    dv_cat.showErrorMessage = True
                    dv_cat.errorStyle = "information"
                    dv_cat.errorTitle = "Category Notice"
                    dv_cat.error = f"Existing categories: {', '.join(cat_names)}. You can also type a new category name."
                    ws.add_data_validation(dv_cat)
                    dv_cat.add("B2:B500")
            except Exception as e:
                logger.error(f"Error adding category validation dropdown: {e}")

            # Status Dropdown
            try:
                dv_status = DataValidation(type="list", formula1='"active,inactive"', allow_blank=True)
                dv_status.hide_drop_down = None
                dv_status.showInputMessage = True
                dv_status.promptTitle = "Status Selection"
                dv_status.prompt = "Select 'active' or 'inactive'."
                dv_status.showErrorMessage = True
                dv_status.errorStyle = "information"
                dv_status.errorTitle = "Status Selection"
                dv_status.error = "Please select either 'active' or 'inactive'."
                ws.add_data_validation(dv_status)
                status_idx = headers.index('Status') + 1 if 'Status' in headers else 6
                status_col = openpyxl.utils.get_column_letter(status_idx) if hasattr(openpyxl.utils, 'get_column_letter') else chr(64 + status_idx)
                dv_status.add(f"{status_col}2:{status_col}500")
            except Exception as e:
                logger.error(f"Error adding status validation dropdown: {e}")

            # Membership Plan Name Dropdown (Col G when in paid_plan mode)
            if 'Membership Plan Name' in headers:
                try:
                    plans = get_tenant_query(MembershipPlan).filter_by(tenant_id=tenant_id, status='active').all()
                    plan_names = [p.name or p.plan_name for p in plans if (p.name or p.plan_name)]
                    plan_names = [p.replace('"', '').strip() for p in plan_names if p]
                    if plan_names:
                        plan_formula = f'"{",".join(plan_names)}"'
                        dv_plan = DataValidation(type="list", formula1=plan_formula, allow_blank=True)
                        dv_plan.hide_drop_down = None
                        dv_plan.showInputMessage = True
                        dv_plan.promptTitle = "Membership Plan Selection"
                        dv_plan.prompt = f"Select an active plan from dropdown: {', '.join(plan_names)}"
                        dv_plan.showErrorMessage = True
                        dv_plan.errorStyle = "information"
                        dv_plan.errorTitle = "Membership Plan Selection"
                        dv_plan.error = f"Active membership plans: {', '.join(plan_names)}"
                        ws.add_data_validation(dv_plan)
                        plan_idx = headers.index('Membership Plan Name') + 1
                        plan_col = openpyxl.utils.get_column_letter(plan_idx) if hasattr(openpyxl.utils, 'get_column_letter') else chr(64 + plan_idx)
                        dv_plan.add(f"{plan_col}2:{plan_col}500")
                except Exception as e:
                    logger.error(f"Error adding membership plan validation dropdown: {e}")

            g.use_master_db = was_master

        elif module_name == 'customers':
            try:
                dv_gender = DataValidation(type="list", formula1='"Female,Male,Other"', allow_blank=True)
                dv_gender.hide_drop_down = None
                dv_gender.showInputMessage = True
                dv_gender.promptTitle = "Gender Selection"
                dv_gender.prompt = "Select Female, Male, or Other."
                dv_gender.showErrorMessage = True
                dv_gender.errorStyle = "information"
                dv_gender.errorTitle = "Gender Selection"
                dv_gender.error = "Please select Female, Male, or Other."
                ws.add_data_validation(dv_gender)
                dv_gender.add("E2:E500")
            except Exception as e:
                logger.error(f"Error adding customer gender dropdown: {e}")

        elif module_name in ['employees', 'products']:
            try:
                dv_status = DataValidation(type="list", formula1='"active,inactive"', allow_blank=True)
                dv_status.hide_drop_down = None
                dv_status.showInputMessage = True
                dv_status.promptTitle = "Status Selection"
                dv_status.prompt = "Select 'active' or 'inactive'."
                dv_status.showErrorMessage = True
                dv_status.errorStyle = "information"
                dv_status.errorTitle = "Status Selection"
                dv_status.error = "Please select either 'active' or 'inactive'."
                ws.add_data_validation(dv_status)
                status_idx = headers.index('Status') + 1 if 'Status' in headers else len(headers)
                status_col = openpyxl.utils.get_column_letter(status_idx) if hasattr(openpyxl.utils, 'get_column_letter') else chr(64 + status_idx)
                dv_status.add(f"{status_col}2:{status_col}500")
            except Exception as e:
                logger.error(f"Error adding status dropdown: {e}")

        # Save to bytes
        file_stream = BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)
        
        return file_stream.getvalue()
    except Exception as e:
        logger.error(f"Error generating Excel template for {module_name}: {str(e)}")
        raise

def get_sample_data(module_name, membership_mode="paid_plan"):
    """Get sample data for the template"""
    if module_name == 'customers':
        return [
            {
                'first_name': 'John',
                'last_name': 'Doe',
                'phone': '9876543210',
                'email': 'john.doe@example.com',
                'gender': 'Male',
                'date_of_birth': '1990-01-15',
                'address': '123 Main Street',
                'notes': 'VIP Customer',
                'spot': 'A1'
            }
        ]
    elif module_name == 'employees':
        return [
            {
                'first_name': 'Jane',
                'last_name': 'Smith',
                'phone': '9876543211',
                'specialization': 'Hair Stylist',
                'role': 'Stylist',
                'salary': '25000',
                'commission_percentage': '10',
                'joining_date': '2023-01-01',
                'status': 'active'
            }
        ]
    elif module_name == 'services':
        sample = {
            'name': 'Haircut',
            'category_name': 'Hair Services',
            'price': '500',
            'duration_minutes': '30',
            'description': 'Basic haircut',
            'status': 'active'
        }
        if membership_mode in ['paid_plan', 'discount_plan']:
            sample['membership_plan_name'] = 'SUPER PLAN'
            sample['membership_discount_percentage'] = '10'
            sample['membership_discount_amount'] = '50'
        return [sample]
    elif module_name == 'products':
        return [
            {
                'name': 'Shampoo',
                'category': 'Hair Care',
                'sku': 'SHM001',
                'barcode': '1234567890123',
                'cost_price': '150',
                'selling_price': '300',
                'mrp': '350',
                'stock_quantity': '50',
                'low_stock_threshold': '10',
                'status': 'active'
            }
        ]
    return []

def parse_excel_file(file_data, module_name):
    """
    Parse uploaded Excel file and return structured data with validation
    Returns: (success, data, errors)
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(BytesIO(file_data))
        ws = wb.active
        
        # Get headers from first row
        headers = []
        for cell in ws[1]:
            if cell.value:
                headers.append(str(cell.value).strip())
        
        # Validate headers
        field_mapping = FIELD_MAPPINGS.get(module_name, {})
        required = REQUIRED_FIELDS.get(module_name, [])
        
        # Check for missing required headers
        missing_headers = [h for h in required if h not in headers]
        if missing_headers:
            return False, [], f"Missing required columns: {', '.join(missing_headers)}"
        
        # Check for invalid headers
        valid_headers = set(field_mapping.keys())
        invalid_headers = [h for h in headers if h not in valid_headers]
        if invalid_headers:
            return False, [], f"Invalid columns found: {', '.join(invalid_headers)}"
        
        # Parse data rows
        data = []
        errors = []
        
        for row_num, row in enumerate(ws.iter_rows(min_row=2), 2):
            row_data = {}
            row_errors = []
            
            for col_num, header in enumerate(headers, 1):
                cell_value = ws.cell(row=row_num, column=col_num).value
                field_name = field_mapping[header]
                
                # Convert cell value to appropriate type
                row_data[field_name] = convert_cell_value(cell_value, field_name)
            
            # Validate required fields
            for required_header in required:
                field_name = field_mapping[required_header]
                if not row_data.get(field_name):
                    row_errors.append(f"{required_header} is required")
            
            if row_errors:
                errors.append({
                    'row': row_num,
                    'errors': row_errors,
                    'data': row_data
                })
            else:
                data.append(row_data)
        
        return True, data, errors
        
    except Exception as e:
        logger.error(f"Error parsing Excel file: {str(e)}")
        return False, [], f"Error parsing Excel file: {str(e)}"

def convert_cell_value(value, field_name):
    """Convert Excel cell value to appropriate Python type"""
    if value is None:
        return None
    
    # Handle different field types
    if field_name in ['salary', 'price', 'cost_price', 'selling_price', 'mrp', 'commission_percentage', 'membership_discount_percentage', 'membership_discount_amount']:
        try:
            return float(str(value))
        except (ValueError, TypeError):
            return 0.0
    elif field_name in ['duration_minutes', 'stock_quantity', 'low_stock_threshold']:
        try:
            return int(str(value))
        except (ValueError, TypeError):
            return 0
    elif field_name in ['date_of_birth', 'joining_date']:
        if isinstance(value, datetime):
            return value.strftime('%Y-%m-%d')
        return str(value)
    else:
        return str(value).strip() if value else None

def validate_customer_data(row_data, tenant_id, branch_id=None):
    """Validate customer data according to existing rules"""
    from app.models.customer import Customer
    from app.database import db
    from flask import g
    
    errors = []
    
    # Validate required fields
    if not row_data.get('first_name'):
        errors.append("First Name is required")
    if not row_data.get('phone'):
        errors.append("Phone Number is required")
    
    # Validate phone format
    phone = row_data.get('phone', '')
    if phone and len(phone) < 10:
        errors.append("Phone Number must be at least 10 digits")
    
    # Validate email format if provided
    email = row_data.get('email')
    if email and '@' not in email:
        errors.append("Invalid Email format")
    
    # Validate date format if provided
    dob = row_data.get('date_of_birth')
    if dob:
        try:
            datetime.strptime(dob, '%Y-%m-%d')
        except ValueError:
            errors.append("Date of Birth must be in YYYY-MM-DD format")
    
    # Check for duplicate phone in tenant context
    if phone:
        query = Customer.query.filter_by(tenant_id=tenant_id, phone=phone)
        if branch_id:
            query = query.filter_by(branch_id=branch_id)
        existing = query.first()
        if existing:
            errors.append(f"Customer with phone {phone} already exists")
    
    return errors

def validate_employee_data(row_data, tenant_id, branch_id=None):
    """Validate employee data according to existing rules"""
    from app.models.employee import Employee
    from app.database import db
    
    errors = []
    
    # Validate required fields
    if not row_data.get('first_name'):
        errors.append("First Name is required")
    if not row_data.get('phone'):
        errors.append("Phone Number is required")
    if row_data.get('salary') is None:
        errors.append("Salary is required")
    
    # Validate phone format
    phone = row_data.get('phone', '')
    if phone and len(phone) < 10:
        errors.append("Phone Number must be at least 10 digits")
    
    # Validate salary
    salary = row_data.get('salary', 0)
    try:
        salary_val = float(salary)
        if salary_val < 0:
            errors.append("Salary must be >= 0")
    except (ValueError, TypeError):
        errors.append("Invalid Salary value")
    
    # Validate commission percentage
    commission = row_data.get('commission_percentage', 0)
    try:
        comm_val = float(commission)
        if comm_val < 0 or comm_val > 100:
            errors.append("Commission % must be between 0 and 100")
    except (ValueError, TypeError):
        errors.append("Invalid Commission % value")
    
    # Validate joining date if provided
    joining_date = row_data.get('joining_date')
    if joining_date:
        try:
            datetime.strptime(joining_date, '%Y-%m-%d')
        except ValueError:
            errors.append("Joining Date must be in YYYY-MM-DD format")
    
    # Check for duplicate phone in branch context
    if phone:
        query = Employee.query.filter_by(tenant_id=tenant_id, phone=phone)
        if branch_id:
            query = query.filter_by(branch_id=branch_id)
        existing = query.first()
        if existing:
            errors.append(f"Employee with phone {phone} already exists")
    
    return errors

def validate_service_data(row_data, tenant_id):
    """Validate service data according to existing rules"""
    from app.models.catalog import Service, ServiceCategory
    from app.database import db
    
    errors = []
    
    # Validate required fields
    if not row_data.get('name'):
        errors.append("Service Name is required")
    if not row_data.get('category_name'):
        errors.append("Category Name is required")
    if row_data.get('price') is None:
        errors.append("Price is required")
    if row_data.get('duration_minutes') is None:
        errors.append("Duration is required")
    
    # Validate price
    price = row_data.get('price', 0)
    try:
        price_val = float(price)
        if price_val < 0:
            errors.append("Price must be >= 0")
    except (ValueError, TypeError):
        errors.append("Invalid Price value")
    
    # Validate duration
    duration = row_data.get('duration_minutes', 0)
    try:
        dur_val = int(duration)
        if dur_val <= 0:
            errors.append("Duration must be a positive integer")
    except (ValueError, TypeError):
        errors.append("Invalid Duration value")
    
    # Validate category exists in tenant context
    category_name = row_data.get('category_name')
    if category_name:
        category = ServiceCategory.query.filter_by(tenant_id=tenant_id, name=category_name).first()
        if not category:
            errors.append(f"Category '{category_name}' does not exist")
        else:
            row_data['category_id'] = category.id
    
    # Check for duplicate service name in tenant context
    service_name = row_data.get('name')
    if service_name:
        existing = Service.query.filter_by(tenant_id=tenant_id, name=service_name).first()
        if existing:
            errors.append(f"Service '{service_name}' already exists")
    
    return errors

def validate_product_data(row_data, tenant_id):
    """Validate product data according to existing rules"""
    from app.models.catalog import Product
    from app.database import db
    
    errors = []
    
    # Validate required fields
    if not row_data.get('name'):
        errors.append("Product Name is required")
    if row_data.get('cost_price') is None:
        errors.append("Cost Price is required")
    if row_data.get('selling_price') is None:
        errors.append("Selling Price is required")
    if row_data.get('mrp') is None:
        errors.append("MRP is required")
    if row_data.get('stock_quantity') is None:
        errors.append("Stock Quantity is required")
    
    # Validate prices
    for price_field in ['cost_price', 'selling_price', 'mrp']:
        price = row_data.get(price_field, 0)
        try:
            price_val = float(price)
            if price_val < 0:
                errors.append(f"{price_field.replace('_', ' ').title()} must be >= 0")
        except (ValueError, TypeError):
            errors.append(f"Invalid {price_field.replace('_', ' ').title()} value")
    
    # Validate stock
    stock = row_data.get('stock_quantity', 0)
    try:
        stock_val = int(stock)
        if stock_val < 0:
            errors.append("Stock Quantity must be >= 0")
    except (ValueError, TypeError):
        errors.append("Invalid Stock Quantity value")
    
    # Validate low stock threshold if provided
    threshold = row_data.get('low_stock_threshold')
    if threshold is not None:
        try:
            thresh_val = int(threshold)
            if thresh_val < 0:
                errors.append("Low Stock Threshold must be >= 0")
        except (ValueError, TypeError):
            errors.append("Invalid Low Stock Threshold value")
    
    # Check for duplicate SKU in tenant context
    sku = row_data.get('sku')
    if sku:
        existing = Product.query.filter_by(tenant_id=tenant_id, sku=sku).first()
        if existing:
            errors.append(f"Product with SKU '{sku}' already exists")
    
    # Check for duplicate barcode in tenant context
    barcode = row_data.get('barcode')
    if barcode:
        existing = Product.query.filter_by(tenant_id=tenant_id, barcode=barcode).first()
        if existing:
            errors.append(f"Product with Barcode '{barcode}' already exists")
    
    return errors